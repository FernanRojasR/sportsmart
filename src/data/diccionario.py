"""Genera docs/02_datos/diccionario_datos.xlsx con el diccionario y perfil de las 12 tablas."""
import pandas as pd, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
B=os.path.abspath(os.path.join(os.path.dirname(__file__),"..",".."))  # raíz del repo
R=os.path.join(B,"data","raw")
D={
"clientes":("Clientes registrados en la tienda",[
 ("cliente_id","ID","Identificador único del cliente (PK)"),("edad","Numérico","Edad en años"),("genero","Nominal","Género declarado: M / F"),
 ("departamento","Nominal","Departamento de residencia"),("fecha_registro","Fecha","Fecha de alta en la plataforma"),
 ("canal_adquisicion","Nominal","Canal por el que llegó el cliente"),("deporte_preferido","Nominal","Deporte declarado en el perfil")]),
"productos":("Catálogo de productos",[
 ("producto_id","ID","Identificador único del producto (PK)"),("nombre","Texto","Nombre comercial"),("categoria","Nominal","Categoría (5 valores)"),
 ("subcategoria","Nominal","Subcategoría (20 valores)"),("marca","Nominal","Marca del producto"),("precio_lista","Numérico","Precio de lista en soles"),
 ("costo","Numérico","Costo unitario en soles"),("talla_aplica","Lógico","Indica si el producto se vende por talla")]),
"pedidos":("Cabecera de pedidos (ventas)",[
 ("pedido_id","ID","Identificador único del pedido (PK)"),("cliente_id","ID","Cliente que compra (FK a clientes)"),
 ("fecha_pedido","Fecha","Fecha y hora de la compra"),("campania_id","ID","Campaña asociada (FK a campanias); vacío = compra sin campaña"),
 ("monto_total","Numérico","Total pagado en soles"),("estado","Nominal","Entregado / Cancelado / En tránsito")]),
"detalle_pedido":("Ítems de cada pedido (ventas a nivel producto)",[
 ("pedido_id","ID","Pedido (FK a pedidos)"),("producto_id","ID","Producto (FK a productos)"),("cantidad","Numérico","Unidades compradas"),
 ("precio_unitario","Numérico","Precio de lista al momento de la compra (S/)"),("talla","Nominal","Talla comprada; vacío si el producto no usa talla"),
 ("descuento","Numérico","Descuento de campaña aplicado (0 a 0.35)"),("importe","Numérico","cantidad × precio_unitario × (1 − descuento)")]),
"pagos":("Pagos de los pedidos",[
 ("pago_id","ID","Identificador del pago (PK)"),("pedido_id","ID","Pedido pagado (FK a pedidos)"),("metodo","Nominal","Tarjeta, Yape, Plin, PagoEfectivo o contraentrega"),
 ("cuotas","Numérico","Número de cuotas (solo tarjeta de crédito)"),("importe","Numérico","Importe cobrado en soles"),("estado_pago","Nominal","Aprobado / Rechazado / Reembolsado")]),
"campanias":("Campañas de marketing",[
 ("campania_id","ID","Identificador de la campaña (PK)"),("nombre","Texto","Nombre de la campaña"),("fecha_inicio","Fecha","Inicio de vigencia"),
 ("fecha_fin","Fecha","Fin de vigencia"),("canal","Nominal","Canal principal de difusión"),("descuento","Numérico","Descuento ofrecido (fracción)")]),
"marketing":("Interacción de clientes con campañas",[
 ("campania_id","ID","Campaña (FK a campanias)"),("cliente_id","ID","Cliente impactado (FK a clientes)"),
 ("abrio","Lógico","Abrió el mensaje de la campaña"),("hizo_clic","Lógico","Hizo clic en el mensaje")]),
"navegacion":("Actividad web/app agregada por cliente y mes",[
 ("cliente_id","ID","Cliente (FK a clientes)"),("mes","Fecha","Mes de la actividad (AAAA-MM)"),("sesiones","Numérico","Sesiones en el mes"),
 ("paginas_vistas","Numérico","Páginas vistas en el mes"),("minutos_totales","Numérico","Tiempo total de navegación"),
 ("dispositivo","Nominal","App / Web móvil / Web desktop"),("carritos_abandonados","Numérico","Sesiones con carrito no comprado")]),
"atencion":("Tickets de atención al cliente",[
 ("ticket_id","ID","Identificador del ticket (PK)"),("cliente_id","ID","Cliente (FK a clientes)"),("pedido_id","ID","Pedido relacionado (FK a pedidos)"),
 ("tipo","Nominal","Consulta / Reclamo"),("tema","Nominal","Tema del ticket"),("canal","Nominal","WhatsApp, email, web chat o teléfono"),
 ("fecha","Fecha","Fecha de apertura"),("horas_resolucion","Numérico","Horas hasta la resolución"),("resuelto","Lógico","Si el ticket fue resuelto")]),
"opiniones":("Reseñas de clientes",[
 ("opinion_id","ID","Identificador de la reseña (PK)"),("pedido_id","ID","Pedido reseñado (FK a pedidos)"),("cliente_id","ID","Autor (FK a clientes)"),
 ("puntuacion","Numérico","Puntuación de 1 a 5"),("comentario","Texto","Comentario libre"),("fecha","Fecha","Fecha de la reseña")]),
"devoluciones":("Devoluciones de productos",[
 ("devolucion_id","ID","Identificador de la devolución (PK)"),("pedido_id","ID","Pedido (FK a pedidos)"),("producto_id","ID","Producto devuelto (FK a productos)"),
 ("talla","Nominal","Talla devuelta; vacío si no aplica"),("motivo","Nominal","Motivo declarado"),("fecha_devolucion","Fecha","Fecha de la devolución"),
 ("monto_reembolso","Numérico","Monto reembolsado en soles")]),
"logistica":("Envíos de los pedidos",[
 ("envio_id","ID","Identificador del envío (PK)"),("pedido_id","ID","Pedido enviado (FK a pedidos)"),("operador","Nominal","Flota propia, Olva, Shalom o Serpost"),
 ("distancia_km","Numérico","Distancia aproximada del envío"),("costo_envio","Numérico","Costo cobrado al cliente (0 si compra > S/ 199)"),
 ("fecha_envio","Fecha","Fecha de despacho"),("dias_entrega","Numérico","Días reales de entrega; vacío si no se entregó"),
 ("dias_prometidos","Numérico","Plazo prometido (3 Lima/Callao, 6 provincias)"),("entrega_a_tiempo","Lógico","dias_entrega ≤ dias_prometidos")]),
}
hdr=Font(name="Arial",bold=True,color="FFFFFF"); fill=PatternFill("solid",fgColor="14213D"); base=Font(name="Arial",size=10)
thin=Side(style="thin",color="D3CFC3"); brd=Border(bottom=thin)
wb=Workbook(); ws=wb.active; ws.title="Diccionario"
cols=["Tabla","Campo","Tipo","Descripción","Ejemplo","% nulos","Valores únicos"]
ws.append(cols)
res=Workbook().active
tabla_rows=[]
for t,(desc,campos) in D.items():
    df=pd.read_csv(os.path.join(R,f"{t}.csv"))
    assert [c for c,_,_ in campos]==list(df.columns), t
    tabla_rows.append((t,desc,len(df),len(df.columns)))
    for c,tipo,d in campos:
        s=df[c].replace(r"^\s*$",pd.NA,regex=True) if df[c].dtype==object else df[c]
        ej=s.dropna().iloc[0]
        ws.append([t,c,tipo,d,str(ej),round(float(s.isna().mean()),4),int(s.nunique())])
for cell in ws[1]: cell.font=hdr; cell.fill=fill; cell.alignment=Alignment(vertical="center")
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.font=base; cell.border=brd; cell.alignment=Alignment(vertical="top",wrap_text=True)
    row[5].number_format="0.0%"; row[6].number_format="#,##0"
for i,w in enumerate([16,20,11,52,30,9,13],1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
n=ws.max_row
ws.cell(n+2,1,"Fuente: perfil calculado con pandas sobre data/raw/*.csv (dataset sintético SportSmart, semilla 42). % nulos incluye celdas vacías.").font=Font(name="Arial",italic=True,size=9)
t2=wb.create_sheet("Tablas"); t2.append(["Tabla","Descripción","Filas","Columnas"])
for r in tabla_rows: t2.append(list(r))
t2.append(["Total","",f"=SUM(C2:C{len(tabla_rows)+1})",f"=SUM(D2:D{len(tabla_rows)+1})"])
for cell in t2[1]: cell.font=hdr; cell.fill=fill
for row in t2.iter_rows(min_row=2):
    for cell in row: cell.font=base
    row[2].number_format="#,##0"
for c in t2[t2.max_row]: c.font=Font(name="Arial",size=10,bold=True)
for i,w in enumerate([18,46,12,10],1): t2.column_dimensions[get_column_letter(i)].width=w
os.makedirs(os.path.join(B,"docs","02_datos"),exist_ok=True)
out=os.path.join(B,"docs","02_datos","diccionario_datos.xlsx"); wb.save(out); print(out, n-1, "campos")
